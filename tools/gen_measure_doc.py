"""Generate Measure_Checks.docx — full enumerated list of every Measure check.

Pulls dict-driven checks live from the code (onpage_audit.CHECKS,
audit.SEO_CHECKS) so those stay exact; the row-emitter modules
(_row("Label", ...)) are listed from their known labels. Run:

    python tools/gen_measure_doc.py
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pipeline.scanner import onpage_audit, audit

SEV_COLOR = {
    "error": RGBColor(0xC0, 0x39, 0x2B),
    "warn":  RGBColor(0xB9, 0x77, 0x0E),
    "info":  RGBColor(0x2E, 0x86, 0xC1),
    "ok":    RGBColor(0x1E, 0x8A, 0x4C),
    "":      RGBColor(0x33, 0x33, 0x33),
}


def dfs_onpage() -> list[tuple[str, str]]:
    """DataForSEO on-page CHECKS, deduped by label (meta-refresh appears twice)."""
    seen, out = set(), []
    for _flag, (label, _why, _fix, sev, _bw) in onpage_audit.CHECKS.items():
        if label in seen:
            continue
        seen.add(label)
        out.append((label, sev))
    return out


# Our own row-emitter modules — label lists (severity varies per page, left blank).
OUR = {
    "Our on-page SEO (entry page) — pipeline/scanner/audit.py": [
        (c[0] if isinstance(c, (tuple, list)) else c, "") for c in audit.SEO_CHECKS
    ],
    "Our on-page deep (single page) — pipeline/scanner/onpage.py": [
        (c, "") for c in [
            "Charset", "Doctype", "Single title", "Single meta description",
            "Mixed content", "External link safety", "Render-blocking scripts",
            "Image dimensions", "Deprecated HTML", "DOM size", "Link volume",
            "hreflang", "Semantic <main>", "URL length", "URL underscores",
            "URL case", "URL parameters", "Subheadings (H2)", "Heading order",
            "Meta refresh", "Single canonical", "Placeholder text", "Flash",
            "Iframe count", "Inline styles", "Empty links", "Legacy meta keywords",
            "Apple touch icon",
        ]
    ],
    "Technical — pipeline/scanner/extra_checks.py": [
        (c, "") for c in [
            "HTTPS", "Mobile viewport", "Language declared", "Open Graph tags",
            "Twitter/X card", "Favicon", "Rendering (crawler-visible content)",
            "XML sitemap", "Structured data found",
        ]
    ],
    "Video + internal links — pipeline/scanner/extra_checks.py": [
        (c, "") for c in [
            "Video snippets (VideoObject)", "Internal links",
            "Anchor text quality", "Broken internal link",
        ]
    ],
    "E-E-A-T — pipeline/scanner/eeat.py": [
        (c, "") for c in [
            "Author / expertise", "Contact / trust", "Policy links",
            "Social proof", "Authoritativeness", "Citation-ready formatting",
        ]
    ],
    "AEO (answer engines) — pipeline/scanner/audit.py": [
        (c, "") for c in [
            "AI crawler allowed (robots.txt)", "robots.txt present",
            "LocalBusiness schema", "Answer-first structure",
        ]
    ],
    "Performance (CrUX field data) — pipeline/scanner/audit.py": [
        (c, "") for c in [
            "LCP (Largest Contentful Paint)", "INP (Interaction to Next Paint)",
            "CLS (Cumulative Layout Shift)",
        ]
    ],
    "Schema / JSON-LD — pipeline/scanner/schema_check.py": [
        (c, "") for c in [
            "Valid JSON-LD", "Schema @type", "Schema types",
        ]
    ],
    "Content info-gain — pipeline/scanner/content.py": [
        (c, "") for c in [
            "Content depth", "Original data", "Comprehensiveness",
        ]
    ],
    "Validation — pipeline/scanner/validate.py": [
        (c, "") for c in [
            "Sitemap valid", "Sitemap index", "Sitemap size",
            "Sitemap URL scheme", "hreflang set", "hreflang href",
            "hreflang x-default",
        ]
    ],
    "Source code (read-only, GitHub repo) — pipeline/scanner/source_audit.py": [
        (c, "") for c in [
            "Source: framework", "Source: rendering (SSR/CSR)",
            "Source: robots.txt in repo", "Source: sitemap.xml in repo",
        ]
    ],
    "Local / Google Business Profile — pipeline/scanner/business_data.py": [
        (c, "") for c in [
            "Google Business Profile found", "Reviews", "Primary category",
            "NAP (name/address/phone)", "Claimed",
        ]
    ],
    "Web mentions — pipeline/scanner/mentions.py": [
        (c, "") for c in [
            "Web mentions", "Mention sentiment",
        ]
    ],
}

# DataForSEO rankings/keywords/backlinks tools — dynamic per-keyword labels.
DFS_TOOLS = {
    "DataForSEO — rankings & keywords — pipeline/scanner/dataforseo.py": [
        ("Ranked keywords (keywords the domain ranks for)", ""),
        ("Domain overview (total ranking keywords / traffic)", ""),
        ("SERP rank (position for each seed keyword)", ""),
        ("Search volume (monthly searches per keyword)", ""),
        ("Keyword ideas (expansion terms)", ""),
        ("Keyword suggestions (autocomplete-style)", ""),
        ("Keyword gap (competitor ranks, we don't)", ""),
        ("Competitors (organic rivals)", ""),
        ("Keyword difficulty (0-100 per keyword)", ""),
        ("Search intent (informational/commercial/etc.)", ""),
        ("LLM mentions (cited by AI engines)", ""),
        ("Backlinks + referring domains", ""),
        ("Broken backlinks", ""),
        ("Historical / visibility trend", ""),
    ],
}


def add_group(doc, title, rows, start_n):
    doc.add_heading(f"{title}  ({len(rows)})", level=2)
    n = start_n
    for label, sev in rows:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(label)
        run.font.size = Pt(10.5)
        if sev:
            tag = p.add_run(f"   [{sev.upper()}]")
            tag.font.size = Pt(8)
            tag.font.color.rgb = SEV_COLOR.get(sev, SEV_COLOR[""])
            tag.bold = True
        n += 1
    return n


def main():
    doc = Document()

    title = doc.add_heading("SEO / AEO Pipeline — Measure Stage", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Full check inventory (read-only audit)")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    dfs_op = dfs_onpage()

    # Count everything.
    total = len(dfs_op) + sum(len(v) for v in OUR.values()) + \
        sum(len(v) for v in DFS_TOOLS.values())

    intro = doc.add_paragraph()
    intro.add_run("Total check types: ").bold = True
    intro.add_run(f"{total}")
    doc.add_paragraph(
        "Data source rule: if DataForSEO has the tool, we use it; only where "
        "they have no coverage do we run our own check. On-page overlaps on "
        "purpose — title/H1/canonical run both on the entry page (ours) and "
        "site-wide (DataForSEO)."
    )

    doc.add_heading("On-page checks", level=1)
    add_group(doc, "DataForSEO on-page audit (site-wide)", dfs_op, 0)
    add_group(doc, "Our on-page SEO (entry page) — audit.py",
              OUR["Our on-page SEO (entry page) — pipeline/scanner/audit.py"], 0)
    add_group(doc, "Our on-page deep (single page) — onpage.py",
              OUR["Our on-page deep (single page) — pipeline/scanner/onpage.py"], 0)
    add_group(doc, "Technical — extra_checks.py",
              OUR["Technical — pipeline/scanner/extra_checks.py"], 0)

    doc.add_heading("Off-page, content & local checks", level=1)
    for key in [
        "Video + internal links — pipeline/scanner/extra_checks.py",
        "E-E-A-T — pipeline/scanner/eeat.py",
        "AEO (answer engines) — pipeline/scanner/audit.py",
        "Performance (CrUX field data) — pipeline/scanner/audit.py",
        "Schema / JSON-LD — pipeline/scanner/schema_check.py",
        "Content info-gain — pipeline/scanner/content.py",
        "Validation — pipeline/scanner/validate.py",
        "Source code (read-only, GitHub repo) — pipeline/scanner/source_audit.py",
        "Local / Google Business Profile — pipeline/scanner/business_data.py",
        "Web mentions — pipeline/scanner/mentions.py",
    ]:
        add_group(doc, key.split(" — ")[0], OUR[key], 0)

    doc.add_heading("DataForSEO rankings & keywords (off-site data)", level=1)
    for key, rows in DFS_TOOLS.items():
        add_group(doc, key.split(" — ")[0], rows, 0)

    out = "/Users/both/seo_agent/Measure_Checks.docx"
    doc.save(out)
    print(f"saved {out}  ({total} checks)")


if __name__ == "__main__":
    main()
